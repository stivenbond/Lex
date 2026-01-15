import docx
import openpyxl
import sqlite3

def parse_docx_table(file_path):
    doc = docx.Document(file_path)
    data = []
    for table in doc.tables:
        for row in table.rows:
            data.append([cell.text.strip() for cell in row.cells])
    return data

def parse_xlsx_table(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    data = []
    for row in ws.iter_rows(values_only=True):
        if any(row):
            data.append([str(cell).strip() if cell is not None else "" for cell in row])
    return data

def import_syllabus(teacher_id, name, data):
    """
    Expects data as a list of rows (lists).
    The first row might be headers.
    We'll assume the first column is ID/No and second is Name for now,
    but in a real app, we might want a mapping UI.
    """
    conn = sqlite3.connect("lex_database.db")
    cursor = conn.cursor()
    
    cursor.execute("INSERT INTO syllabuses (teacher_id, name) VALUES (?, ?)", (teacher_id, name))
    syllabus_id = cursor.lastrowid
    
    for i, row in enumerate(data):
        if i == 0: continue # Skip header
        if len(row) < 2: continue
        
        topic_id_str = row[0]
        topic_name = row[1]
        
        cursor.execute("INSERT INTO topics (syllabus_id, topic_id_string, name, order_index) VALUES (?, ?, ?, ?)",
                       (syllabus_id, topic_id_str, topic_name, i))
    
    conn.commit()
    conn.close()
    return syllabus_id
